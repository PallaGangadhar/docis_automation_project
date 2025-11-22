from flask import Flask, render_template, request, Response, redirect, session, url_for, send_file, jsonify
from flask_socketio import SocketIO, emit
from CASA import *
from Arris import *
from utlity import *
from db import *
from send_mail import send_mail_to

import re
from decorators import login_required
import bcrypt
from datetime import timedelta, datetime
from generate_html import generate_html_file, json_to_html_file
from flask_paginate import Pagination, get_page_args
import random
import string
import json
from docx2pdf import convert
import time
import pythoncom
from docx import Document

per_page = 5
async_mode = None

app = Flask(__name__)
app.secret_key = 'super secret key'
app.config['SECRET_KEY'] = "testkey"

def capitalize_words(s):
    return ' '.join([word.capitalize() for word in s.split()])

app.config['SESSION_TYPE'] = 'filesystem'

def total_time(regression_id):
    curr, conn=db_connection()
    curr.execute(f'''
                                                      
            WITH extracted_numbers AS (
            SELECT 
                CAST(REGEXP_REPLACE(execution_time, '[^0-9.]', '', 'g') AS NUMERIC) AS numeric_value
            FROM 
                regression_logs_details 
            WHERE 
                regression_id = {regression_id}
        )
        SELECT 
            SUM(numeric_value) AS sum_of_numbers
        FROM 
            extracted_numbers;

                                                        ''')
    total_time_for_each_regression=curr.fetchone()[0]
    conn.commit()
    curr.close()
    conn.close()
    return total_time_for_each_regression
    

def remove_space(regression_name):
    return regression_name.replace(' ','_')

def generate_random_alphanumeric(tab_name):
    characters = string.ascii_letters + string.digits
    return f"{tab_name}_{''.join(random.choices(characters, k=10))}"

def get_tc_status(tc_id, reg_id):
    curr, conn=db_connection()
    curr.execute("select status from regression_logs_details where regression_id = %s and testcase_id = %s",(reg_id, tc_id) )
    res = curr.fetchone()
    status = "InProgress"
    if res is None:
        status = "InProgress"
    else:
        status = res[0]
    curr.close()
    conn.close()
    return status


app.jinja_env.filters['total_time'] = total_time
app.jinja_env.filters['remove_space'] = remove_space
app.jinja_env.filters['generate_random_alphanumeric'] = generate_random_alphanumeric
app.jinja_env.filters['get_tc_status'] = get_tc_status

socketio = SocketIO(app, async_mode=async_mode)


def shutdown_server():
    func = request.environ.get('werkzeug.server.shutdown')
    if func is None:
        raise RuntimeError('Not running with the Werkzeug Server')
    func()

# def send_chart_details(data):
#     pass_tc = data.get('pass')
#     fail_tc = data.get('fail')
#     query_string = data.get('query_string')
#     log_id = data.get('log_id')
#     status = 'PASS' if pass_tc == 1 else 'FAIL'
#     reg_id = get_resgression_id(query_string)
#     tab_id  = get_tab_info(query_string)
#     # if log_id:
#     #     reg_id = get_resgression_id_by_log_id(log_id)
#     #     total_rerun_count = data.get('total_rerun_count')
#     # else:
#     #     reg_id = get_resgression_id(query_string)
    
#     # p = f = 0
#     # if log_id:
#     #     p = p + pass_tc
#     #     f = f + fail_tc
#     #     if total_rerun_count:
#     #         total_count = total_rerun_count
#     #     else:
#     #         total_count = p + f
#     # else:
#     #     pass_count, fail_count, total_count,no_run = select_query_to_get_count_details(reg_id)
#     #     p=int(pass_count)+pass_tc
#     #     f=int(fail_count)+fail_tc
#     pass_count, fail_count, total_count,no_run = select_query_to_get_count_details(reg_id)
#     p=int(pass_count)+pass_tc
#     f=int(fail_count)+fail_tc
#     update_regression(pass_tc=pass_tc, fail_tc=fail_tc, r_id=reg_id, log_id=log_id)
#     # tc_no, tc_name, status = get_tc_execution_status(reg_id)
#     tc_execution_count = get_tc_executed_count(reg_id)

#     socketio.emit('charts_details',
#                   {
#                     'pass_tc':p,
#                     'fail_tc':f,
#                     'r_id':reg_id,
#                     'total_count':total_count,
#                     "tab":tab_id,
#                     'query_string': query_string,
#                     'status': status,
#                     'tc_execution_count':tc_execution_count
#                     # 'tc_number': tc_no,
#                     # 'tc_name':tc_name,
#                     # 'status':status
#                     })

def send_chart_details(data):
    pass_tc = data.get('pass')
    fail_tc = data.get('fail')
    query_string = data.get('query_string')
    log_id = data.get('log_id', None)
    status = 'PASS' if pass_tc == 1 else 'FAIL'
    reg_id = get_resgression_id(query_string)
    tab_id  = get_tab_info(query_string)
    rerun = "False"
    existing_status = None
    validate_status = None
    if log_id:
        reg_id = get_resgression_id_by_log_id(log_id)
        total_rerun_count = data.get('total_rerun_count', None)
        rerun = "True"
        existing_status = get_executed_tc_status(log_id)
        validate_status = validate_tc_status(status, existing_status)
    else:
        reg_id = get_resgression_id(query_string)
    
    p = f = 0
    if log_id is not None:
        # add_track_rerun_count(reg_id, total_rerun_count)
       
        pass_count, fail_count, total_count = get_rerun_pass_fail_count_status(reg_id)
        p=int(pass_count)+pass_tc
        f=int(fail_count)+fail_tc
        
        print("Rerun TC - Pass:", p, "Fail:", f, "Total:", total_count)
        update_track_rerun_count(reg_id, p, f)
    else:
        pass_count, fail_count, total_count,no_run = select_query_to_get_count_details(reg_id)
        p=int(pass_count)+pass_tc
        f=int(fail_count)+fail_tc
    
    if not validate_status and log_id is not None:
        update_regression(pass_tc=pass_tc, fail_tc=fail_tc, r_id=reg_id, log_id=log_id, status=status)
    else:
        if log_id is None:
            update_regression(pass_tc=pass_tc, fail_tc=fail_tc, r_id=reg_id, log_id=log_id)
        # tc_no, tc_name, status = get_tc_execution_status(reg_id)
    tc_execution_count = get_tc_executed_count(reg_id)

    socketio.emit('charts_details',
                  {
                    'pass_tc':p,
                    'fail_tc':f,
                    'r_id':reg_id,
                    'total_count':total_count,
                    "tab":tab_id,
                    'query_string': query_string,
                    'status': status,
                    'tc_execution_count':tc_execution_count,
                    "rerun":rerun
                    })

@app.route('/', methods=['GET','POST'])
@login_required
def index():
    total_regression_count = 0
    graph_data={}
    data1=[]
    dates=[]
    
    date_error_message = ""
    from_date = request.args.get('from_date')
    to_date = request.args.get('to_date')

    if from_date == None and to_date == None:
        to_date = datetime.strftime(datetime.now()+timedelta(days=1), "%Y-%m-%d")
        from_date = datetime.now()-timedelta(days=10)
        from_date = datetime.strftime(from_date, "%Y-%m-%d")
    else:
        if from_date != "" and to_date != "":
            to_date=convert_str_to_date(to_date,"%Y-%m-%d")+timedelta(days=1)
            to_date=convert_date_to_str(to_date,"%Y-%m-%d")
        else:
            date_error_message = "Please select end date."
    curr, conn=db_connection()
    curr.execute('SELECT count(regression_logs_details.log_id) from  regression_logs_details')
    total_regression_count=curr.fetchone()
    total_regression_count=total_regression_count[0]
    curr.execute(f"SELECT * FROM devices_details ORDER BY device_id DESC")
    devices_details=curr.fetchall()
    
    curr.execute('SELECT devices_details.device_name, count(regression_logs_details.*) FROM regression_logs_details inner JOIN regression ON regression.regression_id = regression_logs_details.regression_id RIGHT JOIN devices_details ON regression.device_id = devices_details.device_id GROUP BY devices_details.device_name')
    devices_regression_count = curr.fetchall()
    total_regression_devices = len(devices_regression_count)+1
    
    is_get_diff = get_date_diff(from_date, to_date)
    if is_get_diff == ">30 and <=364":
        curr.execute(f"SELECT COUNT(*) AS total_count, TO_CHAR(date_added, 'YYYY-MM') AS month_year FROM regression WHERE date_added BETWEEN '{from_date}' AND '{to_date}' GROUP BY month_year ORDER BY month_year DESC;")
    elif is_get_diff == ">365":
        curr.execute(f"SELECT COUNT(*) AS total_count, TO_CHAR(date_added, 'YYYY') AS year FROM regression WHERE date_added BETWEEN '{from_date}' AND '{to_date}' GROUP BY year ORDER BY year DESC;")
    else:
        curr.execute(f"SELECT count(*),DATE(date_added) as reg_date FROM regression WHERE date_added BETWEEN '{from_date}' AND '{to_date}' GROUP BY DATE(date_added) ORDER BY reg_date DESC")
    regression_graph=curr.fetchall()

    curr.execute(f"SELECT devices_details.device_name, count(regression_logs_details.regression_id) from  regression INNER JOIN regression_logs_details ON regression_logs_details.regression_id = regression.regression_id INNER JOIN devices_details ON regression.device_id = devices_details.device_id AND regression.date_added BETWEEN '{from_date}' AND '{to_date}' GROUP BY devices_details.device_name ORDER BY devices_details.device_name")
    pie_chart_details = curr.fetchall()  
    for i in regression_graph:   
        date = i[1] if is_get_diff == ">30 and <=364" or is_get_diff == ">365" else i[1].strftime("%Y-%m-%d")
        data1.append({'date':date,'count':i[0]})
        dates.append(date)
    graph_data['data']=data1
    pie_chart = []
    for pie in pie_chart_details:
        per = (pie[1]/total_regression_count)*100
        pie_chart.append({'name':pie[0], 'y':round(per,2)})

    conn.commit()
    curr.close()
    conn.close()
    is_data_present=len(regression_graph)
    if is_data_present == 0:
        date_error_message = "No Data Found for the selected date range or for last 10 days."
    return render_template('index.html',total_regression_count=total_regression_count,
                        regression_date_graph=graph_data,devices_details=devices_details,
                        devices_regression_count=devices_regression_count,pie_chart=pie_chart,
                        total_regression_devices=total_regression_devices,is_data_present=is_data_present,
                        date_error_message=date_error_message)
    


# @app.route('/logs', methods=['GET','POST'])
# @login_required
# def logs():
#     if request.method == "POST":
#         global reg_id
        
#         reg_id=add_regression(request)
#         tc = request.form.get('data')
        
#         for tc_name in tc.split(','):
#             eval(tc_name + "()")
#         call_after_execution(reg_id)
#     return render_template('logs.html')

@app.route('/tc_execution', methods=['GET','POST'])
@login_required
def tc_execution():
    if request.method == "POST":
        global reg_id
        # regression_string = request.form.get('random_string')
        add_regression(request)
        random_string = request.form.get('random_string')
        device_type = request.form.get('cmts_type')
        testcase_id = request.form.get('testcase_id')
        for  tc_id in testcase_id.split(','):
            tc_func = get_testcase_function(tc_id)
            eval(f"{tc_func}(tc_id={tc_id},random_string ='{random_string}')")

        # for tc_func, tc_id in zip(tc.split(','), testcase_id.split(',')):
        #     eval(f"{tc_func}(tc_id={tc_id})")

        # RDP_call_after_execution()
        # eval(device_type + "_call_after_execution()")
        # eval(device_type +'_result_summary("' + reg_id + '", "'+random_string+'")')
    devices_details=header()
    curr, conn=db_connection()

    curr.execute(f"SELECT distinct vendor FROM devices_details")
    vendor_details = curr.fetchall()

    conn.commit()
    curr.close()
    conn.close()
    return render_template('tc_execution.html',vendor_details=vendor_details,devices_details=devices_details)


@app.route('/modules', methods=['GET','POST'])
@login_required
def modules():
    devices_details=header()
    curr, conn=db_connection()
    search_module = request.args.get('search_module')
    device_type = request.args.get('device_type_dropdown')
    total_details = get_table_total_details(
        **{'table_name':'modules_details','search_module':search_module,'device_type':device_type}
    )
    page, _, _ = get_page_args(page_parameter='page', per_page_parameter='per_page')
    offset = (page - 1) * per_page
    offset_details = get_offset_details(
        **{'table_name':'modules_details','offset':offset, 'per_page':per_page, 'search_module':search_module,'device_type':device_type}
    )
    pagination = Pagination(page=page, per_page=per_page, total=total_details, css_framework='bootstrap4')
    conn.commit()
    curr.close()
    conn.close()
    return render_template('modules.html',
                           pagination=pagination,
                           offset_details=offset_details,
                           devices_details=devices_details
                           )

@app.route('/devices', methods=['GET','POST'])
@login_required
def devices():
    devices_details=header()
    curr, conn=db_connection()
    search_devices = request.args.get('search_device')
    total_details = get_table_total_details(
        **{'table_name':'devices_details','search_devices':search_devices}
    )
    page, _, _ = get_page_args(page_parameter='page', per_page_parameter='per_page')
    offset = (page - 1) * per_page
    offset_details = get_offset_details(
        **{'table_name':'devices_details','offset':offset, 'per_page':per_page, 'search_devices':search_devices}
    )
    pagination = Pagination(page=page, per_page=per_page, total=total_details, css_framework='bootstrap4')
    conn.commit()
    curr.close()
    conn.close()
    return render_template('device_details.html',
                           pagination=pagination,
                           offset_details=offset_details,
                           devices_details=devices_details
                           )

@app.route('/testcase_details', methods=['GET','POST'])
@login_required
def testcase_details():
    devices_details=header()
    tc_name = request.args.get('search_tc')
    module_type = request.args.get('module_type_dropdown')
    device_type = request.args.get('device_type_dropdown')
    
    total_testcase_details = get_table_total_details(
        **{'table_name':'testcase_details','tc_name':tc_name,'module_type':module_type, 'device_type':device_type}
    )
    page, _, _ = get_page_args(page_parameter='page', per_page_parameter='per_page')
    offset = (page - 1) * per_page
    testcase_offset_details = get_offset_details(
        **{
            'table_name':'testcase_details',
           'offset':offset,
           'per_page':per_page,
           'tc_name':tc_name,
           'module_type':module_type,
           'device_type':device_type
        }
    )
    pagination = Pagination(page=page, per_page=per_page, total=total_testcase_details, css_framework='bootstrap4')

    return render_template(
        'testcase_details.html',
        testcase_offset_details=testcase_offset_details,
        devices_details=devices_details,
        pagination=pagination
        )


@app.route('/add_modules_details', methods=['GET','POST'])
@login_required
def add_modules():
    devices_details=header()
    curr, conn=db_connection()
    if request.method == "POST":
        module = str(request.form.get('module')).strip(' ')
        device_id = str(request.form.get('device_id')).strip(' ')
        curr.execute('''INSERT INTO modules_details(device_id,module_name) VALUES (%s,%s)''',(device_id,module) )
    curr.execute(f"SELECT * FROM devices_details ORDER BY device_id DESC")
    devices_details=curr.fetchall()
    conn.commit()
    curr.close()
    conn.close()
    
    return render_template('add_modules.html',devices_details=devices_details)

@app.route('/add_device_details', methods=['GET','POST'])
@login_required
def add_device_details():
    devices_details=header()
    error_message=None
    if request.method == "POST":
        device_name = str(request.form.get('device_name')).strip(' ')
        device_ip = str(request.form.get('device_ip')).strip(' ')
        model = str(request.form.get('model')).strip(' ')
        vendor = str(request.form.get('vendor')).strip(' ')
        curr, conn=db_connection()
        
        try:
            curr.execute('''INSERT INTO devices_details(device_name,ip,model, vendor) VALUES (%s,%s,%s,%s)''',(device_name,device_ip, model, vendor) )
        except Exception as e:
            pat = re.search("DETAIL:.*", str(e))
            if pat!= None:
                error_message=pat.group(0)
            else:
                error_message = str(e)
            
        conn.commit()
        curr.close()
        conn.close()
    return render_template('add_device_details.html',error_message=error_message,devices_details=devices_details)


@app.route('/add_testcase_details', methods=['GET','POST'])
@login_required
def add_testcase_details():
    devices_details=header()
    curr, conn=db_connection()
    if request.method == "POST":
        module_id = str(request.form.get('module_id')).strip(' ')
        testcase_number = str(request.form.get('testcase_number')).strip(' ')
        testcase_name = str(request.form.get('testcase_name')).strip(' ')
        testcase_function = str(request.form.get('testcase_function')).strip(' ')
        testcase_reference=str(request.form.get('testcase_reference')).strip(' ')
        testcase_reference_details=str(request.form.get('testcase_reference_details')).strip(' ')
        testcase_reference_details = testcase_reference_details.replace(' ','@')
        curr.execute('''INSERT INTO testcase_details(modules_id,testcase_number,testcase_name,testcase_function,testcase_reference,testcase_reference_details) VALUES (%s,%s,%s,%s,%s,%s)''',(module_id,testcase_number,testcase_name,testcase_function,testcase_reference,testcase_reference_details) )
        
    curr.execute(f"SELECT * FROM modules_details ORDER BY modules_id DESC")
    modules_details=curr.fetchall()
    conn.commit()
    curr.close()
    conn.close()
    return render_template('add_testcase_details.html',modules_details=modules_details,devices_details=devices_details)


@app.route('/send_message', methods=['POST'])
def send_message():
    message = request.json.get('data')
    tab = request.json.get('query_string')
    rerun = request.json.get('rerun')
    if tab:
        socketio.emit('message', {'msg': f'{message}', 'tab': tab, 'rerun': rerun})
        return {'status': 'success', 'message': message}
    return {'status': 'error', 'message': 'User not connected'}, 400
 


@app.route("/charts", methods=['GET','POST'])
@socketio.event
def charts():    
    if request.method == "POST":
        data=request.json
        send_chart_details(data)
        return Response({'msg':"Hi"})
    else:
        return render_template('charts.html')

@app.route("/disconnect", methods=['GET','POST'])
@socketio.on('disconnect')
def test_disconnect():
    emit('client disconnected',request.sid)


@app.route("/stop", methods=['GET','POST'])
@socketio.on('disconnect')
def stop():
    if request.method == "POST":
        shutdown_server()
    return Response({'msg':''})


@app.route("/add_regression_logs", methods=['GET','POST'])
def add_regression_logs():
    if request.method == "POST":
        response=request.json
        add_regression_details(response)
    return Response({'msg':''})


@app.route("/view_regression_details", methods=['GET','POST'])
@login_required
def view_regression_details():
    curr,conn=db_connection()
    curr.execute(f"SELECT * FROM devices_details ORDER BY device_id DESC")
    devices_details=curr.fetchall()
    cmts_type = request.args.get('cmts_type_dropdown')
    search_reg = request.args.get('search_reg')

    total_regression_details = get_table_total_details(
        **{'table_name':'regression','cmts_type':cmts_type, 'search_reg':search_reg}
    )
    page, _, _ = get_page_args(page_parameter='page', per_page_parameter='per_page')
    offset = (page - 1) * per_page
    offset_reg_details = get_offset_details(
        **{'table_name':'regression','offset':offset, 'per_page':per_page,'cmts_type':cmts_type, 'search_reg':search_reg}
    )
    pagination = Pagination(page=page, per_page=per_page, total=total_regression_details, css_framework='bootstrap4',href=f"/view_regression_details?page={{0}}")
    # pagination = Pagination(page=page, per_page=per_page, total=total_details, css_framework='bootstrap4',href=f"/get_tc_execution_status?testcase_id={','.join(testcase_id)}&active_tab={active_tab}&tab_name={tab_name}&random_string={random_string}&page={{0}}")

    curr.close()
    conn.close()
    return render_template('regression_details.html', devices_details=devices_details,
                           offset_reg_details=offset_reg_details, pagination=pagination)

@login_required
@app.route("/view_tc_logs_details/<int:reg_id>", methods=['GET','POST'])
def view_tc_logs_details(reg_id):
    devices_details=header()
    curr,conn=db_connection()
    curr.execute(f'SELECT tc.testcase_number, tc.testcase_name, tc.testcase_reference, reg.failed_in, reg.execution_time, reg.status, reg.tc_logs_path, tc.testcase_id, reg.log_id FROM testcase_details tc FULL JOIN regression_logs_details reg ON tc.testcase_id = reg.testcase_id WHERE reg.regression_id = {reg_id} order by status asc')
    tc_logs_details=curr.fetchall()
    curr.execute(f"SELECT regression_name, date_added, summary_path FROM regression WHERE regression_id={reg_id}")
    data=curr.fetchone()
    conn.commit()
    curr.close()
    conn.close()
    date_time = data[1]
    dt_obj = date_time
    dt_obj = date_time.replace(microsecond=0)
    formatted_date = dt_obj.strftime("%Y-%m-%d_%H:%M:%S")
    
    return render_template(
                        'view_tc_logs_details.html',
                        tc_logs_details=tc_logs_details,
                        summary_path=data[2],
                        reg_id=reg_id,
                        devices_details=devices_details,
                        regression_name=data[0],
                        date_added=formatted_date
                    )


@app.route('/delete_regression/<int:id>', methods=['GET','POST'])
def delete_regression(id):
    if request.method == "POST":
        curr,conn=db_connection()
        curr.execute(f'DELETE FROM regression_logs_details WHERE regression_id={id}')
        curr.execute(f'DELETE FROM regression WHERE regression_id={id}')
        conn.commit()
        curr.close()
        conn.close()
        return redirect("/view_regression_details")


@app.route('/send_mail/<int:reg_id>', methods=['GET','POST'])
def send_mail(reg_id):
    if request.method == "POST":
        email= request.form.get('emails')
        
        curr,conn=db_connection()
        curr.execute(f'SELECT summary_path FROM regression WHERE regression_id={reg_id}')
        summary_path = curr.fetchone()
        send_mail_to(summary_path[0])
        conn.commit()
        curr.close()
        conn.close()
        return redirect("/view_regression_details")

@app.route('/delete_selected_regression', methods=['GET','POST'])
def delete_selected_regression():
    tc_ids = request.form.get('data')
    tc_ids=tc_ids.split(',')
    curr,conn=db_connection()
    for ids in tc_ids:
        print("Tc id's ===== ", ids)
        curr.execute(f'DELETE FROM regression_logs_details WHERE regression_id={ids}')
        curr.execute(f'DELETE FROM regression WHERE regression_id={ids}')
    conn.commit()
    curr.close()
    conn.close()
    
    return redirect("/view_regression_details")

@app.route('/delete_all_regressions', methods=['GET','POST'])
def delete_all_regressions():
    curr,conn=db_connection()
    curr.execute(f'DELETE FROM regression_logs_details')
    curr.execute(f'DELETE FROM regression ')
    conn.commit()
    curr.close()
    conn.close()
    
    return redirect("/view_regression_details")

@app.route('/login', methods=['GET','POST'])
def login():
    msg = ''
    if request.method == 'POST' and 'username' in request.form and 'password' in request.form:
        username = request.form['username']
        password = request.form['password']
        curr,conn=db_connection()
        curr.execute('SELECT * FROM user_info WHERE username = %s', (username, ))
        account = curr.fetchone()
        if account != None:
            password = password.encode('utf-8') 
            account_password = account[1].encode('utf-8')
            if password == account_password:
                session['loggedin'] = True
                session['id'] = account[0]
                session['username'] = account[1]
                msg = 'Logged in successfully !'
                return redirect(url_for('index'))
            else:
                msg = 'Incorrect username / password !'
        else:
            msg = 'Incorrect username / password !'

    return render_template('login.html',msg = msg)


@app.route('/register', methods =['GET', 'POST'])
def register():
    msg = ''
    if request.method == 'POST' and 'username' in request.form and 'password' in request.form :
        username = request.form['username']
        password = request.form['password']
        curr,conn=db_connection()
        curr.execute('SELECT * FROM user_info WHERE username = %s', (username, ))
        account = curr.fetchone()
        if account:
            msg = 'Account already exists !'
        elif not re.match(r'[A-Za-z0-9]+', username):
            msg = 'Usernamconverted_to_bytese must contain only characters and numbers !'
        elif not username or not password:
            msg = 'Please fill out the form !'
        else:
            converted_to_bytes = password.encode('utf-8') 
            salt = bcrypt.gensalt() 
            encrypted_password = bcrypt.hashpw(converted_to_bytes, salt)
            curr.execute('''INSERT INTO user_info(username, password) VALUES (%s, %s)''', (username, encrypted_password, ))

            conn.commit()
            msg = 'You have successfully registered !'
            return redirect(url_for('login'))
    elif request.method == 'POST':
        msg = 'Please fill out the form !'
    return render_template('register.html', msg = msg)

 
@app.route('/logout')
def logout():
    session.pop('loggedin', None)
    session.pop('id', None)
    session.pop('username', None)
    return redirect(url_for('login'))


@app.route('/edit_device_details/<int:device_id>', methods=['GET','POST'])
@login_required
def edit_device_details(device_id):
    devices_details=header()
    curr, conn=db_connection()
    error_message = None
    update_device_details=curr.execute(f"SELECT * FROM devices_details WHERE device_id={device_id}")
    update_device_details=curr.fetchone()
    if request.method == "POST":
        device_name = str(request.form.get('device_name'))
        device_ip = str(request.form.get('device_ip'))
        model = str(request.form.get('model'))
        vendor = str(request.form.get('vendor'))
        try:
            curr.execute(f'UPDATE devices_details SET device_name=%s,ip=%s,model=%s,vendor=%s WHERE device_id = %s',(device_name,device_ip, model, vendor,device_id) )
            conn.commit()
            curr.close()
            conn.close()
        except Exception as e:
            pat = re.search("DETAIL:.*", str(e))
            if pat!= None:
                error_message=pat.group(0)
            else:
                error_message = str(e)
        if error_message != None:
            return render_template('edit_device_details.html',update_device_details = update_device_details,error_message=error_message,devices_details=devices_details)
        else:
            return redirect(url_for('devices'))
      
    conn.commit()
    curr.close()
    conn.close()
        
    return render_template('edit_device_details.html',update_device_details = update_device_details,error_message=error_message,devices_details=devices_details)

@app.route('/edit_testcase_details/<int:testcase_id>', methods=['GET','POST'])
@login_required
def edit_testcase_details(testcase_id):
    devices_details=header()
    curr, conn=db_connection()
    error_message = None
    update_testcase_details=curr.execute(f"SELECT * FROM testcase_details WHERE testcase_id={testcase_id}")
    update_testcase_details=curr.fetchone()
    
    curr.execute(f"SELECT modules_id, module_name from modules_details WHERE modules_id = {update_testcase_details[1]}")
    selected_module = curr.fetchone()
    
    curr.execute(f"SELECT * FROM modules_details WHERE modules_id <> {update_testcase_details[1]}")
    modules_details = curr.fetchall()
    if request.method == "POST":
        testcase_number = str(request.form.get('testcase_number'))
        testcase_name = str(request.form.get('testcase_name'))
        testcase_function = str(request.form.get('testcase_function'))
        testcase_reference = str(request.form.get('testcase_reference'))
        module_id = str(request.form.get('module_id'))
        try:
            curr.execute(f'UPDATE testcase_details SET modules_id = %s, testcase_number=%s,testcase_name=%s,testcase_function=%s,testcase_reference=%s WHERE testcase_id = %s',(module_id, testcase_number,testcase_name, testcase_function,testcase_reference,testcase_id) )
            conn.commit()
            curr.close()
            conn.close()
        except Exception as e:
            pat = re.search("DETAIL:.*", str(e))
            if pat!= None:
                error_message=pat.group(0)
            else:
                error_message = str(e)
        if error_message != None:
            return render_template('edit_testcase_details.html',update_testcase_details = update_testcase_details,error_message=error_message,devices_details=devices_details, selected_module=selected_module, modules_details=modules_details)
        else:
            return redirect(url_for('testcase_details'))
      
    conn.commit()
    curr.close()
    conn.close()
        
    return render_template('edit_testcase_details.html',update_testcase_details = update_testcase_details,error_message=error_message,devices_details=devices_details, selected_module=selected_module, modules_details=modules_details)


@app.route('/edit_modules_details/<int:module_id>', methods=['GET','POST'])
@login_required
def edit_module_details(module_id):
    devices_details=header()
    curr, conn=db_connection()
    update_module_details=curr.execute(f"SELECT * FROM modules_details WHERE modules_id={module_id}")
    update_module_details=curr.fetchone()
    error_message = None

    if request.method == "POST":
        module_name = str(request.form.get('module'))
        try:
            curr.execute(f'UPDATE modules_details SET module_name=%s WHERE modules_id = %s',(module_name,module_id) )
            conn.commit()
            curr.close()
            conn.close()
        except Exception as e:
            pat = re.search("DETAIL:.*", str(e))
            if pat!= None:
                error_message=pat.group(0)
            else:
                error_message = str(e)
        return redirect(url_for('modules'))
      
    conn.commit()
    curr.close()
    conn.close()
        
    return render_template('edit_modules_details.html',update_module_details = update_module_details,error_message=error_message,devices_details=devices_details)

@app.route('/get_device_details_from_modules',  methods=['GET','POST'])
def get_device_details_from_modules():
    if request.method == 'POST':
        module_id = request.form.get('data')
        curr, conn=db_connection()
        curr.execute(f"SELECT devices_details.device_name FROM devices_details,modules_details where modules_details.device_id=devices_details.device_id and modules_details.modules_id={module_id}")
        module_device_details=curr.fetchone()
        conn.commit()
        curr.close()
        conn.close()
        return render_template('get_device_details_from_modules.html',module_device_details=module_device_details)

@app.route('/delete_testcase/<int:id>', methods=['GET','POST'])
def delete_testcase(id):
    if request.method == "POST":
        curr,conn=db_connection()
        curr.execute(f'DELETE FROM regression_logs_details WHERE testcase_id={id}')
        conn.commit()
        curr.execute(f'DELETE FROM testcase_details WHERE testcase_id={id}')
        time.sleep(1)
        curr.close()
        conn.close()
        return redirect("/testcase_details")

def header():
    curr, conn=db_connection()
    curr.execute(f"SELECT * FROM devices_details ORDER BY device_id DESC")
    devices_details=curr.fetchall()
    conn.commit()
    curr.close()
    conn.close()
    return devices_details

@app.route('/show_details_mapped_to_devices',  methods=['GET','POST'])
def show_details_mapped_to_devices():
    if request.method == 'POST':
        device_id = request.form.get('data')
        curr, conn=db_connection()
        curr.execute(f"select modules_details.module_name from devices_details, modules_details where modules_details.device_id=devices_details.device_id and devices_details.device_id={device_id}")
        module_device_details=curr.fetchall()
        curr.execute(f"select testcase_details.testcase_name from devices_details, testcase_details,modules_details where testcase_details.modules_id=modules_details.modules_id and modules_details.device_id=devices_details.device_id and devices_details.device_id={device_id}")
        testcase_details=curr.fetchall()
        conn.commit()
        curr.close()
        conn.close()
        return render_template('show_details_mapped_to_devices.html',module_device_details=module_device_details,testcase_details=testcase_details)

@app.route('/delete_device/<int:id>', methods=['GET','POST'])
def delete_device(id):
    if request.method == "POST":
        curr,conn=db_connection()
        curr.execute(f'delete from devices_details where device_id={id}')
        conn.commit()
        curr.close()
        conn.close()
        return redirect("/devices")   
    
@app.route('/show_details_mapped_to_modules',  methods=['GET','POST'])
def show_details_mapped_to_modules():
    if request.method == 'POST':
        module_id = request.form.get('data')
        curr, conn=db_connection()
        curr.execute(f"select testcase_details.testcase_name from  testcase_details where modules_id={module_id}")
        testcase_details=curr.fetchall()
        conn.commit()
        curr.close()
        conn.close()
        return render_template('show_details_mapped_to_modules.html',testcase_details=testcase_details)


@app.route('/delete_module/<int:id>', methods=['GET','POST'])
def delete_module(id):
    if request.method == "POST":
        curr,conn=db_connection()
        curr.execute(f'delete from modules_details where modules_id={id}')
        conn.commit()
        curr.close()
        conn.close()
        return redirect("/modules")  
    
@app.route('/get_module_details',  methods=['GET','POST'])
def get_module_details():
    if request.method == 'POST':
        device_id = request.form.get('data')
        curr, conn=db_connection()
        curr.execute(f"select modules_id, module_name from modules_details where device_id={device_id}")
        module_device_details=curr.fetchall()
        conn.commit()
        curr.close()
        conn.close()
        return render_template('get_module_details.html',module_device_details=module_device_details)

@app.route('/get_selected_module_details',  methods=['GET','POST'])
def get_selected_module_details():
    module_device_details = module_details = ""
    if request.method == 'POST':
        module_id = request.form.get('module_id')
        device_id = request.form.get('device_id')
        curr, conn=db_connection()
        if module_id != "" and device_id != "":
            curr.execute(f"select modules_id, module_name from modules_details where modules_id={module_id}")
            module_details=curr.fetchone()
            curr.execute(f"select modules_id, module_name from modules_details where device_id={device_id} and modules_id != {module_id}")
            module_device_details=curr.fetchall()

        elif module_id == "" and device_id != "":
            curr.execute(f"select modules_id, module_name from modules_details where device_id={device_id}")
            module_device_details=curr.fetchall()
        conn.commit()
        curr.close()
        conn.close()
        
        return render_template('get_selected_module_details.html',module_device_details=module_device_details,
                               module_details = module_details)

@app.route('/generate_html/<int:reg_id>', methods=['GET', 'POST'])
def generate_html(reg_id):
    curr, conn=db_connection()
    output_file_path = 'static\\html_logs\\ganga.html'
    input_file = "C:\\Users\\GANGADHAR PALLA\\Documents\\CMTS_Automation_Project\\CMTS_New_UI_dev\\ganga.txt"
    # print("input file ", input_file)
    # generate_html_file(input_file, output_file_path)
    curr.execute(f"SELECT tc.testcase_number, tc.testcase_name, logs.status, logs.failed_in, logs.execution_time FROM regression_logs_details logs join testcase_details tc on logs.testcase_id = tc.testcase_id and logs.regression_id={reg_id}")
    tc_summary_details=curr.fetchall()
    columns = [desc[0] for desc in curr.description]

    # Convert to list of dicts
    result = [dict(zip(columns, row)) for row in tc_summary_details]

    # Convert to JSON string
    json_data = json.dumps(result, indent=4)
    

    output_file_path = json_to_html_file(json_data, output_file_path)
    curr.close()
    conn.close()
    
    doc = Document("docs_template\\Report_Template.docx")
    find_replace_docx(doc, f"docs_summary\\Report_Summary_{reg_id}.docx",json.loads(json_data))
    time.sleep(2)
    pythoncom.CoInitialize()   # Initialize COM
    convert(f"docs_summary\\Report_Summary_{reg_id}.docx", f"docs_summary\\Report_Summary_{reg_id}.pdf")
    pythoncom.CoUninitialize() # Optional, clean up
    return send_file(output_file_path)


@app.route('/generate_docs/<int:reg_id>', methods=['GET', 'POST'])
def generate_docs(reg_id):
    curr, conn=db_connection()
    output_file_path = 'static\\html_logs\\ganga.html'
    input_file = "C:\\Users\\GANGADHAR PALLA\\Documents\\CMTS_Automation_Project\\CMTS_New_UI_dev\\ganga.txt"
    # print("input file ", input_file)
    # generate_html_file(input_file, output_file_path)
    curr.execute(f"SELECT tc.testcase_number, tc.testcase_name, logs.status, logs.failed_in, logs.execution_time FROM regression_logs_details logs join testcase_details tc on logs.testcase_id = tc.testcase_id and logs.regression_id={reg_id}")
    tc_summary_details=curr.fetchall()
    columns = [desc[0] for desc in curr.description]

    # Convert to list of dicts
    result = [dict(zip(columns, row)) for row in tc_summary_details]

    # Convert to JSON string
    json_data = json.dumps(result, indent=4)
    
    curr.close()
    conn.close()
    
    doc = Document("docs_template\\Report_Template.docx")
    find_replace_docx(doc, f"docs_summary\\Report_Summary_{reg_id}.docx",json.loads(json_data))
    time.sleep(2)
    # pythoncom.CoInitialize()   # Initialize COM
    # convert(f"docs_summary\\Report_Summary_{reg_id}.docx", f"docs_summary\\Report_Summary_{reg_id}.pdf")
    # pythoncom.CoUninitialize() # Optional, clean up
    return send_file(output_file_path)

@app.route("/rerun_tc", methods=['GET', 'POST'])
@app.route("/rerun_tc/<int:log_id>", methods=['GET', 'POST'])
@login_required
def rerun_tc(log_id=None):
    reg_id = request.args.get('reg_id')
    total_rerun_count = request.args.get('total_rerun_count')
    curr, conn=db_connection()
    rerun_tc_details = curr.execute(f''' SELECT testcase_details from track_rerun_count WHERE reg_id = {reg_id}''')
    rerun_tc_details = curr.fetchone()
    t = tuple((x for x in rerun_tc_details[0].split(',')))
    if len(t) == 1:
        t = f"({t[0]})"

    # get_tc_details = curr.execute(f'''SELECT DISTINCT
    #         testcase_number, 
    #         testcase_name,
    #         testcase_id

    #     FROM 
    #         testcase_details

    #     WHERE 
    #         testcase_id in {t} ''')
    # get_tc_details = curr.fetchall()
    curr.close()
    conn.close()

    if request.method == "POST":
        testcase_list = request.form.get('tc_id').split(',')
        logs_id_list = request.form.get('log_id').split(',')
        total_rerun_count = request.form.get('total_rerun_tc')
        random_string = request.args.get('random_string')
        reg_id = get_resgression_id_by_log_id(logs_id_list[0])
        # existing_regression_string = get_regression_random_string(reg_id)
        for tc_id, log_ids in zip(testcase_list, logs_id_list):
            testcase_function = get_testcase_function(tc_id)
            try:
                eval(
                    f"{testcase_function}("
                    f"tc_id='{tc_id}', "
                    f"log_id='{log_ids}', "
                    f"total_rerun_tc={total_rerun_count},"
                    f"reg_id={reg_id},"
                    f"random_string='{random_string}')"
                )
            
            except Exception as e:
                print("Error ===", e)
        reg_id = get_resgression_id_by_log_id(logs_id_list[0])
        delete_track_rerun_tc(reg_id)
    
    return render_template(
        're_run_tc.html',
        total_rerun_count=total_rerun_count
        # get_tc_details=get_tc_details,
        # reg_id=reg_id

        )


@app.route('/get_device_details',  methods=['GET','POST'])
def get_device_details():
    if request.method == 'POST':
        vendor = request.form.get('data')
        curr, conn=db_connection()
        curr.execute("select device_id, device_name from devices_details where vendor = %s",(vendor,) )
        # 'SELECT * FROM user_info WHERE username = %s', (username, )
        device_details=curr.fetchall()
        conn.commit()
        curr.close()
        conn.close()
        return render_template('get_device_details.html',device_details=device_details)
    
@app.route('/get_tab_modules_details',  methods=['GET','POST'])
def get_tab_modules_details():
    if request.method == 'POST':
        device_id = request.form.get('data')
        curr, conn=db_connection()
        curr.execute("select modules_id, module_name from modules_details where device_id = %s",(device_id,) )
        module_tab_details=curr.fetchall()
        conn.commit()
        curr.close()
        conn.close()
        return render_template('get_tab_modules_details.html',module_tab_details=module_tab_details)
    
@app.route('/get_tab_testcase_details',  methods=['GET','POST'])
def get_tab_testcase_details():
    if request.method == 'POST':
        testcase_tab_details = None
        modules_id = request.get_json()['modules']

        if len(modules_id) > 0:
            placeholders = ','.join(['%s'] * len(modules_id)) 
            curr, conn=db_connection()
            query = f"""
                SELECT testcase_id, testcase_name 
                FROM testcase_details 
                WHERE modules_id IN ({placeholders})
            """
            curr.execute(query, modules_id)
            testcase_tab_details=curr.fetchall()
            curr.close()
            conn.close()
        return render_template('get_tab_testcase_details.html',testcase_tab_details=testcase_tab_details)

@app.route('/get_tc_execution_status',  methods=['GET','POST'])
def get_tc_execution_status():

    testcase_id = tuple(request.args.get('testcase_id').split(','))
    if len(testcase_id) == 1:
        testcase_id = f"({testcase_id[0]})"

    active_tab = request.args.get('active_tab')
    tab_name = request.args.get('tab_name')
    random_string = request.args.get('random_string')

    total_details = get_table_total_details(
    **{'filter_for':'get_tc_execution_status','testcase_id':testcase_id, 'random_string':random_string}
    )

    per_page = 3
    page, _, _ = get_page_args(page_parameter='page', per_page_parameter='per_page')
    offset = (page - 1) * per_page
    reg_id = get_resgression_id(random_string)
    offset_details = get_offset_details(
        **{'filter_for':'get_tc_execution_status','testcase_id':testcase_id, 'offset':offset, 'per_page':per_page,'random_string':random_string}
    )

    pagination = Pagination(page=page, per_page=per_page, total=total_details, css_framework='bootstrap4',href=f"/get_tc_execution_status?testcase_id={','.join(testcase_id)}&active_tab={active_tab}&tab_name={tab_name}&random_string={random_string}&page={{0}}")

    return render_template(
                            'get_tc_execution_status.html',
                            testcase_tab_details=offset_details,
                            active_tab = active_tab,
                            tab_name=tab_name,
                            pagination=pagination,
                            reg_id=reg_id,
                            page=page
                        )

@app.route('/track_rerun_tc',  methods=['GET','POST'])
def track_rerun_tc():
    if request.method == 'POST':
        tc_id = request.form.get('tc_id')
        r_id = request.form.get('reg_id')
        total_rerun_count = request.form.get('total_rerun_tc')
        curr, conn = db_connection()
        curr.execute('''INSERT INTO track_rerun_count (reg_id, testcase_details, total_rerun_count) VALUES (%s, %s, %s)''', (r_id, tc_id, total_rerun_count))
    
        conn.commit()
        curr.close()
        conn.close()
        return Response({'msg':''})


if __name__ == '__main__':
    socketio.run(app, debug=True,host="0.0.0.0",port=5000)
